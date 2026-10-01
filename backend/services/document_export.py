from io import BytesIO
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt
from fpdf import FPDF

from .text_utils import sanitize_text


def safe_filename(document_type: str, extension: str) -> str:
    name = re.sub(r"[^A-Za-z0-9]+", "_", document_type.strip()).strip("_")
    return f"{name or 'LegalEase_Document'}.{extension}"


def format_txt(text: str) -> bytes:
    return sanitize_text(text).encode("utf-8")


def format_docx(text: str, doc_type: str) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(doc_type.upper())
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    notice = doc.add_paragraph()
    notice.alignment = WD_ALIGN_PARAGRAPH.CENTER
    nr = notice.add_run("DRAFT - REVIEW BY A QUALIFIED LEGAL PROFESSIONAL")
    nr.bold = True
    nr.font.size = Pt(9)

    for raw in sanitize_text(text).splitlines():
        line = raw.strip()
        if not line:
            continue

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)

        if re.match(r"^(SECTION|ARTICLE)\b", line, re.I) or re.match(r"^\d+[\.)]\s+", line):
            r = p.add_run(line)
            r.bold = True
        elif line.startswith("- "):
            p.style = doc.styles["List Bullet"]
            p.add_run(line[2:])
        else:
            p.add_run(line)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Generated with LegalEase - AI-generated draft; legal review recommended.").font.size = Pt(8)

    output = BytesIO()
    doc.save(output)
    return output.getvalue()


class LegalEasePDF(FPDF):
    def __init__(self, title: str):
        super().__init__()
        self.title_text = title

    def header(self):
        self.set_font("Times", "B", 10)
        self.cell(0, 8, "LegalEase", align="C")
        self.ln(5)

    def footer(self):
        self.set_y(-15)
        self.set_font("Times", "", 8)
        self.cell(0, 8, "LegalEase - AI-generated draft; legal review recommended.", align="C")


def format_pdf(text: str, doc_type: str) -> bytes:
    pdf = LegalEasePDF(doc_type)
    pdf.set_title(doc_type)
    pdf.set_margins(18, 18, 18)
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.add_page()

    pdf.set_font("Times", "B", 16)
    pdf.multi_cell(0, 9, sanitize_text(doc_type).upper(), align="C")
    pdf.set_font("Times", "B", 8)
    pdf.cell(0, 6, "DRAFT - REVIEW BY A QUALIFIED LEGAL PROFESSIONAL", align="C")
    pdf.ln(4)
    pdf.ln(4)

    for raw in sanitize_text(text).splitlines():
        line = raw.strip()
        if not line:
            pdf.ln(3)
            continue

        if re.match(r"^(SECTION|ARTICLE)\b", line, re.I) or re.match(r"^\d+[\.)]\s+", line):
            pdf.set_font("Times", "B", 11)
            pdf.multi_cell(0, 6, line)
        elif line.startswith("- "):
            pdf.set_font("Times", "", 11)
            pdf.multi_cell(0, 6, "- " + line[2:])
        else:
            pdf.set_font("Times", "", 11)
            pdf.multi_cell(0, 6, line)
        pdf.ln(1)

    return bytes(pdf.output())
